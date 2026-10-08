#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl>=3.1", "phonenumbers>=8.13"]
# ///
"""Transforme les cartes scannées en fiches, contacts Apple et liste d'appels.

Usage : traiter.py DOSSIER --outil CHEMIN [--groupe NOM] [--modele sonnet] [--ignorer NOMS] [--sans-contacts]

DOSSIER/scans/ contient les cartes écrites par `cartes-outil` (ID.json + ID.jpg). Claude reçoit le texte lu et la photo.
Relançable sans risque : seules les cartes et les étapes manquantes sont traitées.
Fichiers tenus dans DOSSIER :
  traitees.json     carte -> fiche, ou raison de l'abandon
  personnes.json    les fiches (source de vérité)
  Liste d'appels.xlsx
"""

import argparse
import base64
import fcntl
import json
import os
import re
import shutil
import statistics
import subprocess
import sys
import tempfile
import unicodedata
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import phonenumbers
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

TAILLE_LOT = 12
TAILLE_LOT_MAX = 20  # un lot s'allonge pour ne pas séparer deux cartes liées (recto/verso, collègues)
LOTS_EN_PARALLELE = 3

STATUTS = ["À appeler", "Messagerie", "À rappeler", "Veut échanger", "RDV pris", "Pas intéressé", "Mauvais numéro"]
COLONNES = [  # (titre, largeur)
    ("Statut", 15), ("Prénom", 14), ("Nom", 16), ("Société", 24), ("Poste", 24), ("Secteur", 22),
    ("Mobile", 17), ("Fixe", 17), ("E-mail", 30), ("Ville", 14), ("Site", 24), ("Source", 16),
    ("Scannée le", 12), ("À vérifier", 30), ("Contacts", 13), ("Notes d'appel", 40), ("Carte", 8), ("Réf.", 22),
]

PAYS = {"": "FR", "france": "FR", "belgique": "BE", "suisse": "CH", "luxembourg": "LU", "monaco": "MC",
        "canada": "CA", "quebec": "CA", "espagne": "ES", "italie": "IT", "allemagne": "DE", "royaume-uni": "GB",
        "pays-bas": "NL", "portugal": "PT", "etats-unis": "US", "usa": "US", "maroc": "MA", "tunisie": "TN",
        "algerie": "DZ", "senegal": "SN", "cote d'ivoire": "CI", "reunion": "RE", "martinique": "MQ",
        "guadeloupe": "GP", "irlande": "IE", "autriche": "AT", "danemark": "DK", "suede": "SE", "norvege": "NO",
        "pologne": "PL", "grece": "GR", "israel": "IL", "chine": "CN", "japon": "JP", "inde": "IN", "bresil": "BR"}

GENERIQUES = {"contact", "info", "infos", "accueil", "hello", "bonjour", "commercial", "sales", "admin", "office",
              "direction", "secretariat", "rh", "recrutement", "support", "service", "serviceclient",
              "communication", "marketing", "devis", "agence", "team"}

SYSTEME = "Tu lis des cartes de visite pour en faire des fiches contact. Le texte vient d'un OCR et peut contenir des erreurs."

CONSIGNES = """Voici {n} cartes de visite, dans l'ordre du scan. Pour chaque carte : son identifiant, le texte lu par l'OCR du Mac (lignes de haut en bas ; la taille relative du texte entre crochets, ×2.0 = deux fois plus grand que la moyenne, souvent le nom ou la société ; puis les QR codes et ce que le détecteur automatique a reconnu), et sa photo.

Règles :
1. Une fiche par personne. Une carte de société sans nom de personne donne une fiche avec prénom et nom vides, si elle porte un téléphone ou un e-mail.
2. Recto et verso : quand une carte est l'autre face d'une carte voisine (même graphisme, même société, ou informations complémentaires comme un mobile ou un e-mail), fusionne les deux en une seule fiche qui reprend les informations des deux faces. Fais de même pour deux photos de la même carte (l'une peut être partielle). La fiche liste alors tous les identifiants dans « cartes ». Deux collègues de la même société restent deux fiches distinctes.
3. Mets dans « ignorees » les cartes illisibles et celles qui ne sont pas des cartes de visite.{ignorer}
4. N'invente rien. Recopie numéros, e-mails et sites tels qu'ils sont écrits sur la photo. L'OCR se trompe parfois (« fl » lu « f », « rn » lu « m », O et 0) : la photo fait foi. Si la photo est illisible à cet endroit, garde la lecture de l'OCR, mets a_verifier à true et explique dans « remarque ».
5. Noms en casse normale (DUPONT → Dupont), accents et particules conservés. Société telle qu'elle est écrite, sans la forme juridique (SAS, SARL…) sauf si elle fait partie du nom d'usage.
6. Téléphones : type « mobile », « fixe » ou « fax ». Sans indicatif, ce sont des numéros français.
7. « secteur » : 2 à 5 mots déduits de la carte (ex. « Agence de communication », « Fabricant de mobilier »), vide si on ne peut pas savoir.
8. a_verifier = true dès qu'un champ important (nom, numéro, e-mail) est douteux. « remarque » reste vide sinon.
Chaque carte du lot doit apparaître soit dans une fiche, soit dans « ignorees ».{contexte}"""

CONSIGNE_CONTEXTE = """

La première carte, marquée CONTEXTE, est la dernière du lot précédent : elle est déjà traitée. Ne crée pas de fiche pour elle seule, et ne la mets pas dans « ignorees ». Mais si une carte du lot est son autre face, crée une seule fiche pour les deux, avec les deux identifiants (celui du contexte compris)."""

_S = {"type": "string"}
SCHEMA = {
    "type": "object",
    "properties": {
        "fiches": {"type": "array", "items": {
            "type": "object",
            "properties": {
                "cartes": {"type": "array", "items": _S},
                "prenom": _S, "nom": _S, "societe": _S, "poste": _S, "secteur": _S,
                "telephones": {"type": "array", "items": {
                    "type": "object",
                    "properties": {"numero": _S, "type": {"type": "string", "enum": ["mobile", "fixe", "fax"]}},
                    "required": ["numero", "type"],
                }},
                "emails": {"type": "array", "items": _S},
                "site": _S, "linkedin": _S, "rue": _S, "code_postal": _S, "ville": _S, "pays": _S,
                "a_verifier": {"type": "boolean"}, "remarque": _S,
            },
            "required": ["cartes", "prenom", "nom", "societe", "poste", "secteur", "telephones", "emails", "site",
                         "linkedin", "rue", "code_postal", "ville", "pays", "a_verifier", "remarque"],
        }},
        "ignorees": {"type": "array", "items": {
            "type": "object", "properties": {"carte": _S, "raison": _S}, "required": ["carte", "raison"],
        }},
    },
    "required": ["fiches", "ignorees"],
}


class Conflit(Exception):
    """Un fichier d'état a changé ailleurs (l'autre Mac via iCloud) ou n'est pas encore téléchargé."""


def version(chemin: Path):
    return (chemin.stat().st_mtime_ns, chemin.stat().st_size) if chemin.exists() else None


_versions: dict[Path, object] = {}  # version de chaque fichier d'état au moment où on l'a lu ou écrit


def dans_icloud(chemin: Path) -> bool:
    """Fichier déchargé du Mac par iCloud (Optimiser le stockage) : seul reste un fichier .icloud."""
    if chemin.exists() or not chemin.with_name(f".{chemin.name}.icloud").exists():
        return False
    subprocess.run(["brctl", "download", str(chemin)], capture_output=True)
    return True


def lire_json(chemin: Path, defaut):
    if dans_icloud(chemin):
        raise Conflit(f"{chemin.name} n'est pas encore téléchargé depuis iCloud (téléchargement lancé)")
    _versions[chemin] = version(chemin)
    return json.loads(chemin.read_text()) if chemin.exists() else defaut


def ecrire_json(chemin: Path, donnees) -> None:
    if chemin in _versions and version(chemin) != _versions[chemin]:
        raise Conflit(f"{chemin.name} a été modifié ailleurs pendant le traitement (l'autre Mac ?)")
    tmp = chemin.with_suffix(".tmp")
    tmp.write_text(json.dumps(donnees, ensure_ascii=False, indent=2))
    tmp.replace(chemin)
    _versions[chemin] = version(chemin)


def sans_accents(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn").lower()


# ---------- Lecture par Claude ----------

def decrire_carte(carte: dict) -> str:
    hauteurs = [l["h"] for l in carte["lignes"]] or [1]
    med = statistics.median(hauteurs) or 1
    lignes = [f"[×{l['h'] / med:.1f}] {l['texte']}" for l in carte["lignes"]]
    d = carte.get("detecte", {})
    detecte = "; ".join(f"{k} : {', '.join(v)}" for k, v in
                        [("téléphones", d.get("telephones", [])), ("e-mails", d.get("emails", [])),
                         ("sites", d.get("sites", []))] if v)
    morceaux = [f"=== Carte {carte['id']} ===", *lignes]
    if carte.get("qr"):
        morceaux.append("QR code : " + " | ".join(carte["qr"]))
    if detecte:
        morceaux.append("Détecté : " + detecte)
    return "\n".join(morceaux)


def appeler_claude(cartes: list[dict], modele: str, scans: Path, contexte: dict | None = None,
                   ignorer: str = "") -> dict:
    """Envoie les cartes (texte OCR + photo) à `claude -p` et renvoie les fiches structurées."""
    claude = shutil.which("claude") or str(Path.home() / ".local/bin/claude")
    regle_ignorer = f" Ignore aussi les cartes de : {ignorer} (ce sont nos propres cartes)." if ignorer else ""
    contenu = [{"type": "text", "text": CONSIGNES.format(
        n=len(cartes), ignorer=regle_ignorer, contexte=CONSIGNE_CONTEXTE if contexte else "")}]
    for carte in ([contexte] if contexte else []) + cartes:
        entete = "CONTEXTE (déjà traitée)\n" if carte is contexte else ""
        contenu.append({"type": "text", "text": entete + decrire_carte(carte)})
        photo = scans / carte.get("image", f"{carte['id']}.jpg")
        if photo.exists():
            contenu.append({"type": "image", "source": {"type": "base64", "media_type": "image/jpeg",
                                                        "data": base64.b64encode(photo.read_bytes()).decode()}})
    message = json.dumps({"type": "user", "message": {"role": "user", "content": contenu}})
    # Session isolée : sans outils, hooks, MCP ni CLAUDE.md (40 fois moins de jetons).
    # Les images passent par l'entrée stream-json, qui impose la sortie stream-json.
    cmd = [claude, "-p", "--model", modele, "--input-format", "stream-json", "--output-format", "stream-json",
           "--verbose", "--json-schema", json.dumps(SCHEMA), "--tools", "", "--strict-mcp-config",
           "--no-session-persistence", "--setting-sources", "", "--system-prompt", SYSTEME]
    derniere_erreur = ""
    for _ in range(2):
        r = subprocess.run(cmd, input=message + "\n", capture_output=True, text=True, timeout=900,
                           cwd=tempfile.gettempdir())
        resultat = None
        for ligne in r.stdout.splitlines():
            try:
                evt = json.loads(ligne)
            except json.JSONDecodeError:
                continue
            if evt.get("type") == "result":
                resultat = evt
        if resultat and not resultat.get("is_error") and resultat.get("structured_output"):
            return resultat["structured_output"]
        derniere_erreur = (resultat or {}).get("result") or (r.stderr or r.stdout).strip()[-500:] or "réponse vide"
    raise RuntimeError(derniere_erreur)


# ---------- Contrôles ----------

def distance(a: str, b: str) -> int:
    """Distance de Levenshtein."""
    prec = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        cour = [i]
        for j, cb in enumerate(b, 1):
            cour.append(min(prec[j] + 1, cour[j - 1] + 1, prec[j - 1] + (ca != cb)))
        prec = cour
    return prec[-1]


def verifier(fiche: dict, cartes: list[dict]) -> dict:
    """Normalise les numéros et refuse ce qui n'apparaît pas dans le texte lu."""
    lignes = [l["texte"] for c in cartes for l in c["lignes"]] + [q for c in cartes for q in c.get("qr", [])]
    # Chiffres ligne par ligne : un numéro ne doit pas naître de deux champs voisins (code postal + téléphone).
    chiffres_par_ligne = [re.sub(r"\D", "", l) for l in lignes]
    compact = re.sub(r"\s", "", "\n".join(lignes)).lower()
    problemes = []
    pays = sans_accents(fiche["pays"]).strip()
    region = PAYS.get(pays)
    if region is None:
        region = "FR"
        problemes.append(f"pays « {fiche['pays']} » inconnu : numéros sans indicatif lus comme français")

    tels = []
    for t in fiche["telephones"]:
        try:
            n = phonenumbers.parse(t["numero"], region)
        except phonenumbers.NumberParseException:
            n = None
        if n is None or not phonenumbers.is_valid_number(n):
            problemes.append(f"numéro invalide « {t['numero']} »")
            continue
        if not any(str(n.national_number) in ch for ch in chiffres_par_ligne):
            problemes.append(f"numéro {t['numero']} absent du texte lu")
        type_ = t["type"]
        if type_ != "fax" and phonenumbers.number_type(n) == phonenumbers.PhoneNumberType.MOBILE:
            type_ = "mobile"
        format_ = phonenumbers.PhoneNumberFormat.NATIONAL if n.country_code == 33 \
            else phonenumbers.PhoneNumberFormat.INTERNATIONAL
        tels.append({"numero": phonenumbers.format_number(n, phonenumbers.PhoneNumberFormat.E164),
                     "affichage": phonenumbers.format_number(n, format_), "type": type_})

    emails = []
    lus = [m.lower() for l in lignes for m in re.findall(r"[^\s@:;,]+@[^\s@:;,]+", l)]  # e-mails lus par l'OCR
    for e in fiche["emails"]:
        e = re.sub(r"\s", "", e).lower().removeprefix("mailto:")
        if not re.fullmatch(r"[^@]+@[^@]+\.[a-z]{2,}", e):
            problemes.append(f"e-mail invalide « {e} »")
            continue
        if e not in compact:
            # Claude lit la photo et corrige l'OCR (« forent » → « florent ») : quelques lettres d'écart.
            proche = min(lus, key=lambda x: distance(x, e), default=None)
            if proche is not None and distance(proche, e) <= 3:
                problemes.append(f"e-mail lu « {proche} » par l'OCR, corrigé en « {e} » d'après la photo : à confirmer")
            else:
                problemes.append(f"e-mail {e} absent du texte lu")
        emails.append(e)

    remarque = "; ".join(r for r in [fiche["remarque"], *problemes] if r)
    return dict(fiche, telephones=tels, emails=emails, remarque=remarque,
                a_verifier=bool(fiche["a_verifier"] or problemes))


# ---------- Étapes ----------

def mots(carte: dict) -> set[str]:
    return {m for l in carte["lignes"] for m in re.split(r"[^a-z0-9]+", sans_accents(l["texte"])) if len(m) >= 3}


def liees(a: dict, b: dict) -> bool:
    """Deux cartes qui partagent nom de société, domaine ou adresse : recto/verso, double photo ou collègues."""
    ma, mb = mots(a), mots(b)
    return bool(ma and mb) and len(ma & mb) / len(ma | mb) >= 0.15


def decouper_en_lots(ids: list[str], cartes: dict) -> list[list[str]]:
    lots, lot = [], []
    for i in ids:
        if len(lot) >= TAILLE_LOT_MAX or (len(lot) >= TAILLE_LOT and not liees(cartes[lot[-1]], cartes[i])):
            lots.append(lot)
            lot = []
        lot.append(i)
    return lots + [lot] if lot else lots


def meme_personne(a: dict, b: dict) -> bool:
    """Même personne, vue sur deux cartes ou deux photos. Le standard et contact@ ne suffisent pas."""
    def ident(f):
        return sans_accents(f"{f['prenom']} {f['nom']}").strip()
    def perso(f):
        emails = {e for e in f["emails"] if re.sub(r"[^a-z0-9]", "", e.split("@")[0]) not in GENERIQUES}
        return emails | {t["numero"] for t in f["telephones"] if t["type"] == "mobile"}
    meme_societe = sans_accents(a["societe"]) == sans_accents(b["societe"]) != ""
    if ident(a) and ident(b):
        return ident(a) == ident(b) and (meme_societe or bool(perso(a) & perso(b)))
    if ident(a) or ident(b):
        return False
    return meme_societe  # deux cartes de société sans nom de personne


def meme_identite(a: dict, b: dict) -> bool:
    """Preuve positive qu'il s'agit de la même carte : même nom, ou, si l'une n'a pas de nom, même société."""
    na, nb = sans_accents(f"{a['prenom']} {a['nom']}").strip(), sans_accents(f"{b['prenom']} {b['nom']}").strip()
    if na and nb:
        return na == nb
    return sans_accents(a["societe"]) == sans_accents(b["societe"]) != ""


def fusionner(existante: dict, nouvelle: dict) -> list[str]:
    """Complète une fiche avec une autre lecture de la même personne. Renvoie ce qui a été ajouté."""
    ajouts = []
    for champ in ("prenom", "nom", "societe", "poste", "secteur", "site", "linkedin", "rue", "code_postal", "ville", "pays"):
        if not existante[champ] and nouvelle[champ]:
            existante[champ] = nouvelle[champ]
            ajouts.append(f"{champ.replace('_', ' ')} {nouvelle[champ]}")
    numeros = {t["numero"] for t in existante["telephones"]}
    for t in nouvelle["telephones"]:
        if t["numero"] not in numeros:
            existante["telephones"].append(t)
            ajouts.append(f"{t['type']} {t['affichage']}")
    for e in nouvelle["emails"]:
        if e not in existante["emails"]:
            existante["emails"].append(e)
            ajouts.append(f"e-mail {e}")
    # Les doutes d'une lecture qui n'apporte rien (photo partielle, doublon) ne salissent pas la fiche.
    if ajouts:
        existante["a_verifier"] = existante["a_verifier"] or nouvelle["a_verifier"]
        existante["remarque"] = "; ".join(r for r in [existante["remarque"], nouvelle["remarque"]] if r)
        if existante.get("contact"):  # déjà dans Contacts : l'outil ne modifie pas une fiche existante
            existante["a_reporter"] = existante.get("a_reporter", []) + ajouts
        if existante.get("dans_tableur"):
            existante["a_completer"] = True
    existante["cartes"] += [c for c in nouvelle["cartes"] if c not in existante["cartes"]]
    return ajouts


def extraire(dossier: Path, modele: str, ignorer: str = "") -> int:
    """Lit les cartes nouvelles. Renvoie le nombre de cartes qui restent à lire (lot en échec, carte oubliée)."""
    scans = dossier / "scans"
    traitees = lire_json(dossier / "traitees.json", {})
    personnes = lire_json(dossier / "personnes.json", [])
    # Reprise après interruption, ou synchronisation iCloud partielle : traitees.json et personnes.json
    # doivent concorder. Une carte dans une fiche est traitée ; une carte rattachée à une fiche absente est relue.
    dans_une_fiche = {c: f["cle"] for f in personnes for c in f["cartes"]}
    for c in [c for c, v in traitees.items() if not v.startswith("ignorée") and c not in dans_une_fiche]:
        del traitees[c]
    traitees.update(dans_une_fiche)
    absentes = [f for f in scans.glob(".*.json.icloud") if dans_icloud(scans / f.name[1:-len(".icloud")])]
    if absentes:
        print(f"{len(absentes)} carte(s) encore dans iCloud : téléchargement lancé, lues au prochain lancement.")
    ordre = {c: i for i, c in enumerate(sorted(p.stem for p in scans.glob("*.json")))}
    a_faire = sorted(c for c in ordre if c not in traitees)
    # Une photo déchargée par iCloud : on la télécharge et la carte attend le prochain lancement.
    photos_absentes = [c for c in a_faire if dans_icloud(scans / f"{c}.jpg")]
    if photos_absentes:
        print(f"{len(photos_absentes)} photo(s) encore dans iCloud : téléchargement lancé, cartes lues au prochain lancement.")
        a_faire = [c for c in a_faire if c not in photos_absentes]
    if not a_faire:
        print("Aucune nouvelle carte à lire.")
        return len(photos_absentes)
    cartes = {i: json.loads((scans / f"{i}.json").read_text()) for i in a_faire}
    lots = decouper_en_lots(a_faire, cartes)
    print(f"Lecture de {len(a_faire)} carte(s) par Claude ({modele}), {len(lots)} lot(s)…")

    def voisines(a: dict, b: dict) -> bool:
        """Même société et scannées l'une juste après l'autre : recto et verso."""
        positions = {ordre[c] for c in b["cartes"] if c in ordre}
        return sans_accents(a["societe"]) == sans_accents(b["societe"]) != "" \
            and any(ordre[c] + d in positions for c in a["cartes"] if c in ordre for d in (-1, 1))

    def enregistrer_lot(ids: list[str], res: dict, contexte: str | None) -> str:
        nouvelles = []
        vus = set()
        for fiche in res["fiches"]:
            ids_fiche = [c for c in fiche["cartes"] if (c in ids or c == contexte) and c not in vus]
            if not any(c in ids for c in ids_fiche):  # rien de ce lot (fiche du seul contexte)
                continue
            vus.update(ids_fiche)
            f = verifier(fiche, [cartes[c] for c in ids_fiche])
            premiere = cartes[ids_fiche[0]]
            f.update(cle=ids_fiche[0], cartes=ids_fiche, source=premiere.get("source", ""),
                     scannee_le=premiere.get("scanneeLe", "")[:10], contact=None)
            if any(not (scans / f"{c}.jpg").exists() for c in ids_fiche):
                f["a_verifier"] = True
                f["remarque"] = "; ".join(r for r in [f["remarque"], "photo absente : lue sur le seul texte OCR"] if r)
            if contexte in ids_fiche:
                # Autre face de la dernière carte du lot précédent : on complète sa fiche, mais seulement
                # avec une preuve (même nom, ou même société pour un verso sans nom).
                face = next((p for p in personnes if contexte in p["cartes"]), None)
                if face and not meme_identite(face, f):
                    # Gardée à part, et revérifiée sur ses seules cartes : rien ne vient du contexte sans contrôle.
                    seules = [c for c in ids_fiche if c != contexte]
                    f = verifier(fiche, [cartes[c] for c in seules])
                    f.update(cle=seules[0], cartes=seules, source=cartes[seules[0]].get("source", ""),
                             scannee_le=cartes[seules[0]].get("scanneeLe", "")[:10], contact=None, a_verifier=True,
                             remarque="; ".join(r for r in [f["remarque"], "rapprochée par Claude d'une autre carte "
                                                            "sans preuve d'identité : gardée à part"] if r))
                elif face:
                    f["cartes"] = [c for c in ids_fiche if c != contexte]
                    fusionner(face, f)
                    for c in f["cartes"]:
                        traitees[c] = face["cle"]
                    continue
            nouvelles.append(f)
        ignorees = {i["carte"]: i["raison"] for i in res["ignorees"] if i["carte"] in ids and i["carte"] not in vus}
        vus.discard(contexte)
        fusionnees = 0
        for f in nouvelles:
            nommee = bool(f["prenom"] or f["nom"])
            # Une personne déjà connue (autre lot, scan précédent) : on complète sa fiche.
            double = next((p for p in personnes if meme_personne(p, f)), None)
            if double is None and not nommee:  # verso sans nom, recto déjà enregistré
                double = next((p for p in personnes if voisines(f, p)), None)
            if double is None and nommee:
                # Recto arrivé après son verso sans nom : la fiche nommée absorbe le verso,
                # tant que celui-ci n'est ni dans Contacts ni dans la liste.
                verso = next((p for p in personnes if not (p["prenom"] or p["nom"]) and voisines(f, p)
                              and not p.get("contact") and not p.get("dans_tableur")), None)
                if verso:
                    personnes.remove(verso)
                    fusionner(f, verso)
            if double:
                fusionner(double, f)
                fusionnees += 1
            else:
                personnes.append(f)
            cible = double or f
            for c in cible["cartes"]:
                traitees[c] = cible["cle"]
        for c, raison in ignorees.items():
            traitees[c] = f"ignorée : {raison}"
        ecrire_json(dossier / "personnes.json", personnes)
        ecrire_json(dossier / "traitees.json", traitees)
        oubliees = set(ids) - vus - set(ignorees)
        msg = f"  lot de {len(ids)} : {len(nouvelles) - fusionnees} fiche(s), {len(ignorees)} ignorée(s)"
        if fusionnees:
            msg += f", {fusionnees} rattachée(s) à une fiche existante"
        return msg + (f", {len(oubliees)} oubliée(s)" if oubliees else "")

    # Chaque lot reçoit la dernière carte du lot précédent : un recto et son verso séparés par le
    # découpage peuvent ainsi être réunis.
    contextes = [None] + [lot[-1] for lot in lots[:-1]]

    def lire_lot(k: int):
        try:
            ctx = cartes[contextes[k]] if contextes[k] else None
            return appeler_claude([cartes[i] for i in lots[k]], modele, scans, ctx, ignorer), None
        except Exception as e:  # le lot sera repris au prochain lancement
            return None, e

    # Claude lit les lots en parallèle ; on les enregistre dans l'ordre du scan.
    with ThreadPoolExecutor(LOTS_EN_PARALLELE) as pool:
        for k, (res, err) in enumerate(pool.map(lire_lot, range(len(lots)))):
            if err:
                print(f"  lot {lots[k][0]}… en échec : {err}", file=sys.stderr)
            else:
                print(enregistrer_lot(lots[k], res, contextes[k]))
    enregistrees = lire_json(dossier / "traitees.json", {})
    return len(photos_absentes) + sum(1 for i in a_faire if i not in enregistrees)


def ajouter_contacts(dossier: Path, outil: str, groupe: str) -> None:
    personnes = lire_json(dossier / "personnes.json", [])
    a_faire = [p for p in personnes if not p.get("contact") or p["contact"]["statut"] == "erreur"]
    if not a_faire:
        return
    entree = [{"cle": p["cle"], "prenom": p["prenom"], "nom": p["nom"], "societe": p["societe"],
               "poste": p["poste"], "telephones": [{"numero": t["numero"], "type": t["type"]} for t in p["telephones"]],
               "emails": p["emails"], "site": p["site"], "linkedin": p["linkedin"], "rue": p["rue"],
               "codePostal": p["code_postal"], "ville": p["ville"], "pays": p["pays"]} for p in a_faire]
    with tempfile.TemporaryDirectory() as tmp:
        fin, fout = Path(tmp) / "entree.json", Path(tmp) / "sortie.json"
        fin.write_text(json.dumps(entree, ensure_ascii=False))
        r = subprocess.run([outil, "contacts", str(fin), str(fout), "--groupe", groupe], capture_output=True, text=True)
        if r.returncode != 0:
            print(f"Contacts non mis à jour : {r.stderr.strip()}", file=sys.stderr)
            return
        resultats = {x["cle"]: x for x in json.loads(fout.read_text())}
    for p in personnes:
        if p["cle"] in resultats:
            x = resultats[p["cle"]]
            p["contact"] = {"statut": x["statut"], "identifiant": x["identifiant"], "detail": x["detail"]}
    ecrire_json(dossier / "personnes.json", personnes)
    compte = {s: sum(1 for x in resultats.values() if x["statut"] == s) for s in ("cree", "existant", "ambigu", "erreur")}
    print(f"Contacts (groupe « {groupe} ») : {compte['cree']} ajouté(s), {compte['existant']} déjà présent(s)"
          + (f", {compte['ambigu']} ambigu(s) à vérifier" if compte["ambigu"] else "")
          + (f", {compte['erreur']} en erreur" if compte["erreur"] else ""))


LIBELLE_CONTACT = {"cree": "Ajouté", "existant": "Déjà présent", "ambigu": "À vérifier", "erreur": "Erreur"}


def ligne_tableur(p: dict) -> dict:
    mobiles = [t for t in p["telephones"] if t["type"] == "mobile"]
    fixes = [t for t in p["telephones"] if t["type"] == "fixe"]
    return {"Statut": "À appeler", "Prénom": p["prenom"], "Nom": p["nom"], "Société": p["societe"],
            "Poste": p["poste"], "Secteur": p["secteur"], "Mobile": mobiles[0] if mobiles else None,
            "Fixe": fixes[0] if fixes else None, "E-mail": p["emails"][0] if p["emails"] else "",
            "Ville": p["ville"], "Site": p["site"], "Source": p["source"], "Scannée le": p["scannee_le"],
            "À vérifier": "; ".join(r for r in [
                p["remarque"] if p["a_verifier"] else "",
                f"Contacts : {p['contact']['detail']}" if (p.get("contact") or {}).get("statut") == "ambigu" else "",
                "Complétée après coup, à reporter dans Contacts : " + ", ".join(p["a_reporter"])
                if p.get("a_reporter") else "",
            ] if r),
            "Contacts": LIBELLE_CONTACT.get((p.get("contact") or {}).get("statut"), ""), "Réf.": p["cle"]}


def nouveau_classeur() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Appels"
    ws.append([t for t, _ in COLONNES])
    derniere = ws.cell(1, len(COLONNES)).column_letter
    for i, (_, largeur) in enumerate(COLONNES, start=1):
        c = ws.cell(1, i)
        ws.column_dimensions[c.column_letter].width = largeur
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F3E4E")
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{derniere}1"
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUTS) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("A2:A5000")
    couleurs = {"RDV pris": "C6EFCE", "Veut échanger": "C6EFCE", "À rappeler": "FFEB9C", "Messagerie": "FFEB9C",
                "Pas intéressé": "E7E6E6", "Mauvais numéro": "E7E6E6"}
    for statut, fond in couleurs.items():
        ws.conditional_formatting.add(f"A2:{derniere}5000",
                                      FormulaRule(formula=[f'$A2="{statut}"'], fill=PatternFill("solid", bgColor=fond)))
    return wb


def ecrire_cellule(c, v) -> None:
    if isinstance(v, dict):  # téléphone : cliquable sur l'iPhone
        c.value, c.hyperlink = v["affichage"], f"tel:{v['numero']}"
    else:
        c.value = v


def tenir_tableur(dossier: Path) -> bool:
    """Ajoute les nouvelles fiches à la liste d'appels, sans jamais réécrire ce que l'utilisateur y a saisi.
    Renvoie False si la liste n'a pas pu être mise à jour."""
    chemin = dossier / "Liste d'appels.xlsx"
    if (dossier / f"~${chemin.name}").exists():
        print(f"{chemin.name} est ouvert dans Excel : ferme-le puis lance « cartes traiter ».", file=sys.stderr)
        return False
    if dans_icloud(chemin):
        print(f"{chemin.name} n'est pas encore téléchargé depuis iCloud : relance « cartes traiter » dans un instant.",
              file=sys.stderr)
        return False
    personnes = lire_json(dossier / "personnes.json", [])
    etat_lu = version(chemin)
    nouveau = etat_lu is None  # classeur absent : toutes les fiches y sont (ré)écrites
    wb = nouveau_classeur() if nouveau else load_workbook(chemin)
    ws = wb.active
    # Colonnes repérées par leur titre : on peut en ajouter ou en déplacer.
    col = {ws.cell(1, i).value: i for i in range(1, ws.max_column + 1) if ws.cell(1, i).value}
    if "Réf." not in col:
        print(f"Colonne « Réf. » introuvable dans {chemin.name} : liste non mise à jour.", file=sys.stderr)
        return False
    lignes = {ws.cell(r, col["Réf."]).value: r for r in range(2, ws.max_row + 1)}

    # Dernière ligne réellement remplie (max_row compte aussi les lignes vidées à la main).
    fin = max((r for r in range(2, ws.max_row + 1)
               if any(ws.cell(r, c).value not in (None, "") for c in col.values())), default=1)
    ajoutees = completees = 0
    modifie = nouveau
    for p in personnes:
        if p["cle"] in lignes:  # ligne déjà là (relance après une interruption)
            p["dans_tableur"] = True
        if p.get("dans_tableur") and not nouveau:
            # Une ligne supprimée à la main ne revient pas. Sur une ligne existante, le script ne remplit
            # que des cases vides (fiche complétée par une autre photo) et la colonne Contacts.
            if p["cle"] not in lignes:
                continue
            r = lignes[p["cle"]]
            valeurs = ligne_tableur(p)
            if "Contacts" in col and valeurs["Contacts"] and ws.cell(r, col["Contacts"]).value != valeurs["Contacts"]:
                ws.cell(r, col["Contacts"]).value = valeurs["Contacts"]
                modifie = True
            # Seulement si une autre photo a complété la fiche : une case vidée à la main reste vide sinon.
            if p.pop("a_completer", False):
                for titre, v in valeurs.items():
                    if titre in col and titre not in ("Statut", "Contacts", "À vérifier") and v not in (None, "") \
                            and ws.cell(r, col[titre]).value in (None, ""):
                        ecrire_cellule(ws.cell(r, col[titre]), v)
                        completees += 1
                if p.get("a_reporter") and "À vérifier" in col:
                    message = "Complétée après coup, à reporter dans Contacts : " + ", ".join(p.pop("a_reporter"))
                    p["remarque"] = "; ".join(x for x in [p["remarque"], message] if x)
                    c = ws.cell(r, col["À vérifier"])  # ajouté à la suite de ce qui y est déjà
                    c.value = f"{c.value} ; {message}" if c.value else message
                    c.font = Font(color="C00000")
                modifie = True
            continue
        fin += 1
        r = fin
        for titre, v in ligne_tableur(p).items():
            if titre in col:
                ecrire_cellule(ws.cell(r, col[titre]), v)
        if p["emails"] and "E-mail" in col:
            ws.cell(r, col["E-mail"]).hyperlink = f"mailto:{p['emails'][0]}"
        if "Carte" in col:
            carte = ws.cell(r, col["Carte"])
            carte.value, carte.hyperlink = "voir", f"scans/{p['cartes'][0]}.jpg"
        if "À vérifier" in col and ws.cell(r, col["À vérifier"]).value:
            ws.cell(r, col["À vérifier"]).font = Font(color="C00000")
        if p.get("a_reporter"):  # affiché sur la ligne ; enregistré seulement si le classeur l'est
            p["remarque"] = "; ".join(x for x in [p["remarque"], "à reporter dans Contacts : "
                                                  + ", ".join(p.pop("a_reporter"))] if x)
        p.pop("a_completer", None)
        p["dans_tableur"] = True
        ajoutees += 1
        modifie = True

    if not modifie:
        print(f"Liste d'appels : rien de nouveau → {chemin}")
        return True
    # Écriture dans un fichier temporaire puis remplacement : jamais de classeur à moitié écrit.
    tmp = chemin.with_name(f".{chemin.stem}.tmp.xlsx")
    wb.save(tmp)
    # Le classeur a-t-il changé depuis sa lecture (Excel sur l'iPhone, l'autre Mac, iCloud) ? Contrôle fait
    # juste avant le remplacement.
    if version(chemin) != etat_lu:
        tmp.unlink(missing_ok=True)
        print(f"{chemin.name} a changé pendant le traitement : rien n'est écrasé, relance « cartes traiter ».",
              file=sys.stderr)
        return False
    os.replace(tmp, chemin)
    ecrire_json(dossier / "personnes.json", personnes)
    print(f"Liste d'appels : {ajoutees} ligne(s) ajoutée(s)"
          + (f", {completees} case(s) complétée(s)" if completees else "") + f" → {chemin}")
    return True


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("dossier", type=Path)
    ap.add_argument("--outil", required=True, help="chemin de cartes-outil")
    ap.add_argument("--groupe", default="Cartes de visite", help="groupe Contacts où ranger les nouvelles fiches")
    ap.add_argument("--modele", default="sonnet")
    ap.add_argument("--sans-contacts", action="store_true", help="ne pas toucher aux Contacts")
    ap.add_argument("--changer-de-mac", action="store_true", help="traiter désormais ce dossier depuis ce Mac")
    ap.add_argument("--ignorer", default="", help="noms ou sociétés dont les cartes sont ignorées (vos propres cartes)")
    a = ap.parse_args()

    # Un seul Mac traite ce dossier : iCloud ne coordonne pas deux Mac qui écrivent en même temps.
    mac = subprocess.run(["scutil", "--get", "ComputerName"], capture_output=True, text=True).stdout.strip()
    marque = a.dossier / ".mac-de-traitement"
    if not marque.exists():
        marque.write_text(mac)
    elif marque.read_text().strip() != mac and not a.changer_de_mac:
        sys.exit(f"Ce dossier est traité par le Mac « {marque.read_text().strip()} ». Pour le traiter désormais "
                 f"depuis « {mac} », laisse iCloud finir de synchroniser, puis relance avec --changer-de-mac.")
    elif a.changer_de_mac:
        marque.write_text(mac)

    # Un seul traitement à la fois sur ce Mac.
    verrou = open(a.dossier / ".traitement.lock", "w")
    try:
        fcntl.flock(verrou, fcntl.LOCK_EX | fcntl.LOCK_NB)
    except BlockingIOError:
        sys.exit("Un traitement est déjà en cours sur ce dossier.")

    try:
        restantes = extraire(a.dossier, a.modele, a.ignorer)
        if not a.sans_contacts:
            ajouter_contacts(a.dossier, a.outil, a.groupe)
        liste_a_jour = tenir_tableur(a.dossier)
    except Conflit as e:
        sys.exit(f"{e}. Rien n'est écrasé : relance « cartes traiter » dans un instant.")

    personnes = lire_json(a.dossier / "personnes.json", [])
    a_verifier = sum(1 for p in personnes if p["a_verifier"])
    resume = f"{len(personnes)} fiche(s) au total, dont {a_verifier} à vérifier."
    if restantes:
        resume += f" {restantes} carte(s) pas encore lue(s) : relance « cartes traiter »."
    if not liste_a_jour:
        resume += " Liste d'appels pas à jour : relance « cartes traiter »."
    print(resume)
    subprocess.run(["osascript", "-e", f'display notification "{resume}" with title "Cartes de visite"'],
                   capture_output=True)
    if restantes or not liste_a_jour:
        sys.exit(1)


if __name__ == "__main__":
    main()
